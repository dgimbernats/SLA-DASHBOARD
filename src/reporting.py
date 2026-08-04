"""Streamlit-independent Excel and PDF report generation."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.kpi_engine import backlog_aging, summary_kpis, technician_csat

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def create_excel_report(data: pd.DataFrame) -> bytes:
    """Return a formatted XLSX workbook as bytes."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["KPI", "Value"])
    labels = {
        "total_tickets": "Total tickets",
        "open_backlog": "Open backlog",
        "sla_compliance": "SLA compliance (%)",
        "fcr_rate": "FCR rate (%)",
        "reopen_rate": "Reopen rate (%)",
        "avg_csat": "Average CSAT",
    }
    for key, value in summary_kpis(data).items():
        summary.append([labels[key], value])

    aging = backlog_aging(data)
    _write_rows(workbook.create_sheet("Backlog Aging"), ["Age bucket", "Tickets"], aging.items())
    tech = technician_csat(data)
    _write_rows(
        workbook.create_sheet("Technician CSAT"),
        ["Technician", "Average CSAT", "Responses"],
        tech.itertuples(index=False, name=None),
    )

    export = data.copy()
    for column in export.select_dtypes(include=["datetimetz"]).columns:
        export[column] = export[column].dt.strftime("%Y-%m-%d %H:%M:%S UTC")
    _write_rows(
        workbook.create_sheet("Tickets"),
        list(export.columns),
        export.itertuples(index=False, name=None),
    )

    for sheet in workbook.worksheets:
        _format_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_pdf_report(data: pd.DataFrame) -> bytes:
    """Return a compact executive KPI report as PDF bytes."""
    metrics = summary_kpis(data)
    pdf = FPDF()
    pdf.set_title("SLA Dashboard Report")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, "SLA Dashboard Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 7, f"Generated: {pd.Timestamp.now(tz='UTC'):%Y-%m-%d %H:%M UTC}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    rows = [
        ("Total tickets", metrics["total_tickets"]),
        ("Open backlog", metrics["open_backlog"]),
        ("SLA compliance", f"{metrics['sla_compliance']:.2f}%"),
        ("First-contact resolution", f"{metrics['fcr_rate']:.2f}%"),
        ("Reopen rate", f"{metrics['reopen_rate']:.2f}%"),
        ("Average CSAT", f"{metrics['avg_csat']:.2f} / 5"),
    ]
    _pdf_table(pdf, "Executive KPIs", rows)
    _pdf_table(pdf, "Backlog Aging", list(backlog_aging(data).items()))

    top = technician_csat(data).head(10)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 9, "Top Technician CSAT", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    for row in top.itertuples(index=False):
        label = _pdf_text(str(row.assigned_to))
        pdf.cell(105, 6, label[:60], border=1)
        pdf.cell(35, 6, f"{row.avg_csat:.2f}", border=1, align="R")
        pdf.cell(30, 6, str(row.responses), border=1, align="R", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _write_rows(sheet, headers, rows) -> None:
    sheet.append([_excel_value(value) for value in headers])
    for row in rows:
        sheet.append([_excel_value(value) for value in row])


def _excel_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat()
    return value.item() if hasattr(value, "item") else value


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for index, column in enumerate(sheet.columns, start=1):
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        sheet.column_dimensions[get_column_letter(index)].width = width


def _pdf_table(pdf: FPDF, title: str, rows) -> None:
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 9, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    for label, value in rows:
        pdf.cell(120, 7, _pdf_text(str(label)), border=1)
        pdf.cell(50, 7, _pdf_text(str(value)), border=1, align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)


def _pdf_text(value: str) -> str:
    return value.encode("latin-1", "replace").decode("latin-1")
